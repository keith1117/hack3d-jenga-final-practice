from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageOps


ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "figures"

# Put the QR code and its linked image together for the report.
panel = Image.new("RGB", (1600, 720), "white")
draw = ImageDraw.Draw(panel)
draw.text((40, 10), "Recovered QR code", fill="black", font_size=28)
draw.text((740, 10), "Image linked by the QR code", fill="black", font_size=28)
qr = ImageOps.contain(Image.open(FIG / "qr_code.png").convert("RGB"), (650, 650))
photo = ImageOps.contain(Image.open(ROOT / "clues/game_pieces.png").convert("RGB"), (840, 650))
panel.paste(qr, (20, 55))
panel.paste(photo, (730, 55 + (650-photo.height)//2))
panel.save(FIG / "qr_and_linked_image.png")

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = section.bottom_margin = Inches(0.8)
section.left_margin = section.right_margin = Inches(0.9)
for name in ["Normal", "Title", "Heading 1", "Caption"]:
    style = doc.styles[name]
    style.font.name = "Times New Roman"
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = False
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in list(fonts.attrib):
        if "Theme" in key or "theme" in key:
            del fonts.attrib[key]
    for key in ["ascii", "hAnsi", "eastAsia", "cs"]:
        fonts.set(qn("w:" + key), "Times New Roman")
# Remove the decorative title border inherited from the default template.
for border in doc.styles.element.xpath(".//w:pBdr"):
    border.getparent().remove(border)
doc.styles["Normal"].font.size = Pt(12)
doc.styles["Normal"].paragraph_format.line_spacing = 1.1
doc.styles["Normal"].paragraph_format.space_after = Pt(7)
doc.styles["Title"].font.size = Pt(19)
doc.styles["Title"].font.bold = True
doc.styles["Title"].paragraph_format.space_after = Pt(8)
doc.styles["Heading 1"].font.size = Pt(13)
doc.styles["Heading 1"].font.bold = True
doc.styles["Heading 1"].paragraph_format.space_before = Pt(10)
doc.styles["Heading 1"].paragraph_format.space_after = Pt(5)
doc.styles["Caption"].font.size = Pt(10)
doc.styles["Caption"].font.italic = False
doc.core_properties.title = "2023 Hack3D Jenga Practice"
doc.core_properties.author = "Keith Wang"
doc.core_properties.subject = "CSAW final round practice"


def paragraph(text):
    return doc.add_paragraph(text)


def heading(text):
    doc.add_heading(text, level=1)


def figure(filename, width, caption):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(FIG / filename), width=Inches(width))
    doc.add_paragraph(caption, style="Caption")


doc.add_paragraph("2023 Hack3D Jenga Practice", style="Title")
paragraph("Keith Wang  |  Team No Fix  |  October 14, 2026")
paragraph("My answer to this practice challenge is Monopoly. I found a split web address, a QR code, a banknote marked 500, and a small house model in the supplied STEP file. The banknote and house gave me the clearest connection to the game.")

heading("Starting approach")
paragraph("The problem asks me to examine the earlier Jenga release and work out which board game Nyasbro plans to release next. I started with two possibilities: the STEP file might contain a clue in its text, or the clue might be hidden inside the blocks. I searched the file for readable strings and checked the product names. Most of the strings were normal STEP information or part numbers, so this did not give me an answer.")
paragraph("I then inspected the geometry with Python. The file contains 54 numbered part definitions and 55 assembly instances. Part 27 is used twice. Most parts are blocks about 25 by 15 by 75 mm, but part 55 is much smaller. Thirteen part definitions contain enclosed cavities. These differences gave me a manageable place to start.")

heading("The web address")
paragraph("Inside parts 2 and 3, I found thin letter-shaped cuts. I projected their edges onto a flat view and rotated the drawing so I could read it. The address crosses the end of the two blocks, so I also included the open cuts at that shared edge. Joining the two views gave the following address.")
figure("url_joined.png", 6.6, "Figure 1  Internal lettering from parts 2 and 3")
p = paragraph("https://www.britannica.com/biography/J-P-Morgan")
for run in p.runs:
    run.font.size = Pt(10)
paragraph("The address names J. P. Morgan. I treated it as a clue toward money and business, rather than enough evidence to identify a game on its own.")

doc.add_page_break()
heading("Reconstructing the QR code")
paragraph("Several other blocks contained small square patterns. The nested squares looked like the corner markers of a QR code. I grouped the patterns from parts 15, 16, 21, 22, 27, 28, 33, 34, 39, and 40 into five columns. In the original part coordinates, their internal planes are at y values near 37.4, 52.4, 67.4, 82.4, and 97.4 mm.")
paragraph("I projected the patterns onto the x-z plane and placed the columns beside each other in 25 mm steps. The left two pieces were reversed: I moved the pattern in part 15 down by 75 mm and the pattern in part 16 up by 75 mm. This put the three large QR corner markers in sensible positions.")
paragraph("My first QR image was incomplete because I only drew enclosed cavities. Some of the pattern connects to the outside of a block, and different patches use slightly different depths. I included the flat faces throughout the 0.2 mm thick region around each internal plane and kept the white holes inside the black shapes. OpenCV could then decode the result.")
figure("qr_and_linked_image.png", 6.6, "Figure 2  Reconstructed QR code and the board-game image it links to")
paragraph("The decoded link is:")
p = paragraph("https://drive.google.com/file/d/1INetMAmWu2S0kqHK8zMEVcYfBpoYQ0bC/view?usp=sharing")
for run in p.runs:
    run.font.size = Pt(9)
paragraph("I saved the linked PNG in the project. It shows a mixture of board-game pieces, so it supports the general theme but does not point to only one game. I also sampled the QR pattern into a 45 by 45 grid to remove small drawing seams. The raw drawing and the cleaned image decoded to the same link.")

doc.add_page_break()
heading("The banknote and house")
paragraph("Part 19 contains a long rectangular outline with 500 repeated in three places. After projecting the internal edges, it looks like a simplified banknote. I also isolated part 55 from the CAD assembly. Its front view shows a small house with a roof and chimney. I exported that part as both STEP and STL so it can be inspected separately.")
figure("banknote.png", 6.6, "Figure 3  Banknote outline recovered from part 19")
figure("house_front.png", 2.6, "Figure 4  Front view of the house in part 55")
paragraph("The banknote and house fit Monopoly much more directly than the mixed game-piece photo. Hasbro's Monopoly rules list $500 bills and describe buying houses and hotels. Together with the J. P. Morgan address, these clues led me to choose Monopoly as my final answer. I did not use an official answer key.")

heading("What did not work")
paragraph("Searching the STEP strings did not reveal a readable answer. My first point plots were also hard to interpret because vertices alone did not show the full curves or filled regions. Drawing the actual edge curves made the lettering readable. For the QR code, keeping only enclosed shells left large blank areas, and the first arrangement of parts 15 and 16 misplaced the corner markers. Including the open cuts and correcting that order fixed the decoding.")

heading("Files and reference")
paragraph("I kept the original challenge files, extraction scripts, recovered images, part inventory, QR decoding result, and house models in the project repository. The scripts regenerate the main clue drawings from the supplied STEP file.")
p = paragraph("Reference: Hasbro, Monopoly rules, https://www.hasbro.com/common/instruct/Monopoly_Vintage.pdf. The problem statement and STEP model are the supplied 2023 Hack3D challenge materials.")
for run in p.runs:
    run.font.size = Pt(10)

path = ROOT / "Jenga_Practice_Report.docx"
doc.save(path)
print(path)
